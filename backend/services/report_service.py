"""Reports: PDF and Excel."""
import io
from datetime import datetime
from config.supabase_client import supabase
from fastapi.responses import StreamingResponse


def project_summary_excel():
    """Generate Excel report: project summary (admin)."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment
    except ImportError:
        return None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project Summary"
    ws.append(["Project ID", "Title", "Status", "Start", "End", "Modules", "Tasks"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    projects = supabase.table("projects").select("projectid, title, status, start_date, end_date").execute()
    for p in (projects.data or []):
        mods = supabase.table("modules").select("moduleid", count="exact").eq("projectid", p["projectid"]).execute()
        mod_count = getattr(mods, "count", None) or len(mods.data or [])
        task_count = 0
        if mods.data:
            task_ids = supabase.table("tasks").select("taskid").in_("moduleid", [m["moduleid"] for m in mods.data]).execute()
            task_count = len(task_ids.data or [])
        ws.append([
            str(p.get("projectid", "")),
            p.get("title", ""),
            p.get("status", ""),
            str(p.get("start_date", "")),
            str(p.get("end_date", "")),
            mod_count,
            task_count,
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def project_summary_pdf():
    """Generate PDF report: project summary (admin)."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
    except ImportError:
        return None
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [Paragraph("Project Summary Report", styles["Title"]), Spacer(1, 12)]
    projects = supabase.table("projects").select("projectid, title, status, start_date, end_date").execute()
    data = [["Project", "Title", "Status", "Start", "End"]]
    for p in (projects.data or [])[:50]:
        data.append([
            str(p.get("projectid", ""))[:8],
            (p.get("title") or "")[:30],
            p.get("status", ""),
            str(p.get("start_date", "")),
            str(p.get("end_date", "")),
        ])
    t = Table(data)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.grey), ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke)]))
    story.append(t)
    doc.build(story)
    buf.seek(0)
    return buf


def student_progress_excel():
    """Generate Excel: student progress (tasks by status)."""
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        return None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Progress"
    ws.append(["Student User ID", "Total", "Assigned", "In Progress", "Review", "Approved", "Rejected"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    students = supabase.table("students").select("userid").execute()
    for s in (students.data or []):
        uid = s["userid"]
        r = supabase.table("task_assignments").select("status").eq("student_userid", uid).execute()
        by_status = {}
        for row in (r.data or []):
            st = row.get("status") or "assigned"
            by_status[st] = by_status.get(st, 0) + 1
        ws.append([
            uid,
            len(r.data or []),
            by_status.get("assigned", 0),
            by_status.get("in_progress", 0),
            by_status.get("review", 0),
            by_status.get("approved", 0),
            by_status.get("rejected", 0),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
